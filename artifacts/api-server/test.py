import sys
from docx import Document
from docx.shared import Inches

doc = Document('../../local/Reference files/RO2-P60.8t-R90-FF0.85-26.6LMH-20℃-3year-PV6-2(7E-ULP440-MAX).docx')

header = doc.sections[0].header
if len(header.paragraphs) == 0:
    header.add_paragraph()
header.paragraphs[0].add_run().add_picture('../../artifacts/permionics-bd-suite/public/logo.png', width=Inches(1.5))
doc.save('test_out.docx')
