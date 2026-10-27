import sys
from docx import Document
from docx.shared import Inches

def main():
    if len(sys.argv) < 9:
        print("Usage: formatter.py <input_docx> <logo_path> <proj_name> <date> <company> <designed_by> <element_model> <output_docx>")
        sys.exit(1)

    input_docx = sys.argv[1]
    logo_path = sys.argv[2]
    proj_name = sys.argv[3]
    date_val = sys.argv[4]
    company = sys.argv[5]
    designed_by = sys.argv[6]
    element_model = sys.argv[7]
    output_docx = sys.argv[8]

    try:
        doc = Document(input_docx)
    except Exception as e:
        print(f"Error reading docx: {e}")
        sys.exit(1)

    import re

    def normalize_date(date_str: str) -> str:
        if not date_str:
            return ""
        s = date_str.strip()
        m_ymd = re.match(r"^(\d{4})[/-](\d{1,2})[/-](\d{1,2})$", s)
        if m_ymd:
            return f"{int(m_ymd.group(3)):02d}/{int(m_ymd.group(2)):02d}/{m_ymd.group(1)}"
        m_dmy = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$", s)
        if m_dmy:
            return f"{int(m_dmy.group(1)):02d}/{int(m_dmy.group(2)):02d}/{m_dmy.group(3)}"
        return s

    # Dictionary of replacements for standard labels
    replacements = {
        "Project Name:": proj_name,
        "Date:": normalize_date(date_val),
        "Company:": company,
        "Designed By:": designed_by,
    }

    old_models = set()

    # 1. First Pass: Replace standard labels and extract old element models
    for table in doc.tables:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                text = cell.text.strip()
                
                # Remove "Change Your Design" watermark
                if "Change Your Design" in text:
                    cell.text = text.replace("Change Your Design", "")
                
                # Check for Element Model row to extract all old models across all columns
                if text in ["Element Model", "Elements Model"]:
                    # Iterate through all remaining cells in this row
                    for j in range(i + 1, len(row.cells)):
                        old_model_text = row.cells[j].text.strip()
                        if old_model_text and old_model_text not in ["", "N/A", "None"]:
                            old_models.add(old_model_text)
                            if element_model:
                                row.cells[j].text = element_model
                
                # Check for standard label replacements (where the value is in the next cell)
                for label, new_value in replacements.items():
                    if label in text:
                        if len(text) <= len(label) + 2:
                            if i + 1 < len(row.cells) and new_value:
                                row.cells[i+1].text = new_value

    # 2. Second Pass: Globally replace the old element models in all other tables (e.g. Table 3)
    if element_model and old_models:
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    # Exact match replacement
                    if cell.text.strip() in old_models:
                        cell.text = element_model

    # 3. Handle paragraph-level text ("Change Your Design" and inline labels)
    for p in doc.paragraphs:
        if "Change Your Design" in p.text:
            p.text = p.text.replace("Change Your Design", "")
            
        for label, new_value in replacements.items():
            if label in p.text and new_value:
                # E.g. "Project Name: OldVal" -> "Project Name: NewVal"
                if p.text.strip().startswith(label):
                    p.text = f"{label} {new_value}"

    # 3. Replace Header Logo
    if logo_path and logo_path != "dummy_logo":
        try:
            # We add logo to the first section header
            header = doc.sections[0].header
            # Clear existing paragraphs (this removes old text/logos in the header)
            for p in header.paragraphs:
                p.clear()
            
            if len(header.paragraphs) == 0:
                header.add_paragraph()
                
            run = header.paragraphs[0].add_run()
            run.add_picture(logo_path, width=Inches(1.5))
        except Exception as e:
            print(f"Failed to add logo: {e}")

    doc.save(output_docx)
    print("SUCCESS")

if __name__ == "__main__":
    main()
